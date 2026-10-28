"""Write normalized WEHAGO sheets as ordinary ranges, never Excel Tables."""
import argparse
import datetime as dt
import json
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel


def write(data, destination, company, date):
    destination = Path(destination)
    if destination.exists():
        raise FileExistsError(f'Output already exists: {destination}')
    if not data:
        raise ValueError('No sheets')
    wb = Workbook()
    wb.remove(wb.active)
    for d in data:
        headers, rows = d['headers'], d['rows']
        if not headers or any(len(r) != len(headers) for r in rows):
            raise ValueError('Missing header or inconsistent row width')
        if d['name'] in wb.sheetnames:
            raise ValueError('Duplicate sheet name')
        ws = wb.create_sheet(d['name'])
        ws.sheet_view.showGridLines = False
        for cell, value in [('A2', f"{company} {d['name']} — {date}"),
                            ('A3', f"출처: {d.get('source', '')} | {d.get('scope', '')}")]:
            ws[cell] = value
            ws[cell].data_type = 's'
        ws['A2'].font = Font(size=14, bold=True, color='18344F')
        numeric, dates = set(d.get('numeric', [])), set(d.get('dates', []))
        if numeric & dates:
            raise ValueError('Numeric and date columns overlap')
        for ri, row in enumerate([headers] + rows, 5):
            for ci, value in enumerate(row):
                if ri > 5 and value is not None and ci in dates:
                    value = from_excel(value) if isinstance(value, (int, float)) else dt.datetime.fromisoformat(value)
                if ri > 5 and ci in numeric and value is not None:
                    if isinstance(value, bool) or not isinstance(value, (int, float)):
                        raise ValueError(f'Non-numeric value at {d["name"]} {ri}:{ci+1}')
                cell = ws.cell(ri, ci + 1, value)
                if isinstance(value, str):
                    cell.data_type = 's'
                cell.font = Font(name='맑은 고딕', size=10)
                cell.alignment = Alignment(vertical='center', horizontal='right' if ci in numeric and ri > 5 else 'left')
                if ri == 5:
                    cell.fill = PatternFill('solid', fgColor='24476A')
                    cell.font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
                    cell.border = Border(bottom=Side(style='thin', color='BAC5CF'))
                else:
                    if ri % 2 == 0:
                        cell.fill = PatternFill('solid', fgColor='F2F5F8')
                    cell.number_format = ('yyyy-mm-dd' if ci in dates else
                        d.get('number_format', '#,##0.########;[Red](#,##0.########);0') if ci in numeric else '@')
            ws.row_dimensions[ri].height = 21
        for ci in range(len(headers)):
            widths = d.get('widths', [])
            ws.column_dimensions[get_column_letter(ci + 1)].width = widths[ci] if ci < len(widths) else 20
        ws.row_dimensions[5].height = 28
        ws.freeze_panes = 'A6'
        ws.auto_filter.ref = f'A5:{get_column_letter(len(headers))}{5 + len(rows)}'
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    checked = load_workbook(destination)
    for d, ws in zip(data, checked):
        assert not ws.tables, 'Excel Table unexpectedly created'
        assert ws.auto_filter.ref and ws.freeze_panes == 'A6'
        assert ws.max_row == 5 + len(d['rows'])
    checked.close()


if __name__ == '__main__':
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('input')
    p.add_argument('output')
    p.add_argument('--company', required=True)
    p.add_argument('--date', required=True)
    args = p.parse_args()
    dt.date.fromisoformat(args.date)
    write(json.loads(Path(args.input).read_text(encoding='utf-8-sig')), args.output, args.company, args.date)
    print(args.output)
